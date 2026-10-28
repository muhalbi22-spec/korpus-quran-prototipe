"""Build Hikma-formatted .docx files from v2_source.md using the journal's own template.

Usage: build_hikma.py TEMPLATE SOURCE_MD OUT_DOCX VARIANT [COMMENTS_JSON APPENDIX_MD]
VARIANT: named | anon
"""
import copy, json, re, sys
import docx
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

tpl, src, out, variant = sys.argv[1:5]
comments_path = sys.argv[5] if len(sys.argv) > 5 else None
appendix_path = sys.argv[6] if len(sys.argv) > 6 else None

AI_STATEMENT = (
    "[AUTHOR: replace this paragraph with an accurate statement after you have rewritten and verified "
    "the manuscript yourself. Suggested wording, only if it matches what actually happened: \"In preparing "
    "this article, the author used Claude (Anthropic), a generative AI system, to check the dataset against "
    "the Quranic Arabic Corpus, to extract and tabulate translation renderings, and to draft and revise parts "
    "of the text. The author rewrote the text, verified the data, the exegetical sources and the references, "
    "and takes full responsibility for the content; the results and analysis are the product of the author's "
    "own critical judgement.\" Do not keep any claim of verification that you have not carried out.]"
)
SELFREF = ("Fathurrahman, M. A. (2025). Analisis terjemahan istilah teologis pada Al-Qur'an ke dalam bahasa Inggris: "
           "Studi korpus terhadap kata taqwā, raḥmat, kāfir, 'adzāb dan rizq [An analysis of the English translation "
           "of theological terms in the Qur'an: A corpus study of taqwā, raḥmat, kāfir, ʿadhāb and rizq] "
           "[Master's thesis, Universitas Negeri Yogyakarta].")

text = open(src, encoding="utf-8").read()
subs = {
    "named": {"{{SELFCITE}}": "Fathurrahman, 2025", "{{SELFCITE_NARR}}": "Fathurrahman (2025)", "{{SELFREF}}": SELFREF},
    "anon": {"{{SELFCITE}}": "Author, 2025", "{{SELFCITE_NARR}}": "Author (2025)", "{{SELFREF}}": ""},
}[variant]
subs.update({"{{AI_STATEMENT}}": AI_STATEMENT, "{{HASSANEIN_INTRO}}": "", "{{HASSANEIN_REF}}": ""})
for k, v in subs.items():
    text = text.replace(k, v)

d = docx.Document(tpl)
body = d.element.body
sectPr = body.find(qn("w:sectPr"))
for el in list(body):
    if el is not sectPr:
        body.remove(el)

# header: author name / anonymised
for sec in d.sections:
    for hdr in (sec.header, sec.even_page_header, sec.first_page_header):
        for p in hdr.paragraphs:
            for r in p.runs:
                r.text = r.text.replace("Author’s name", "Fathurrahman" if variant == "named" else "Author")
                r.text = r.text.replace("Title in the original language", "Between gloss and faith")

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def add_runs(p, s, size=None, bold=None):
    runs = []
    for part in INLINE.split(s):
        if not part:
            continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*"):
            r = p.add_run(part[1:-1]); r.italic = True
        else:
            r = p.add_run(part)
        if size:
            r.font.size = Pt(size)
        if bold is not None and not part.startswith("**"):
            r.bold = bold
        runs.append(r)
    return runs


def para(s, style="normal", **kw):
    p = d.add_paragraph(style=style)
    add_runs(p, s, **kw)
    return p


def labelled(label, s):
    p = d.add_paragraph(style="Resumen_Abstract")
    r = p.add_run(label); r.bold = True
    add_runs(p, " " + s)
    return p


def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "000000")
        borders.append(el)
    tcPr.append(borders)


def add_table(rows):
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    t = d.add_table(rows=len(rows), cols=ncol)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader"); hdr.set(qn("w:val"), "true"); trPr.append(hdr)
    for row in t.rows:
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); row._tr.get_or_add_trPr().append(cs)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); set_cell_border(c)
            p = c.paragraphs[0]
            pf = p.paragraph_format
            pf.space_before = Pt(0); pf.space_after = Pt(0)
            pf.first_line_indent = Pt(0); pf.left_indent = Pt(0); pf.line_spacing = 1.0
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs(p, val, size=9, bold=True if i == 0 else None)
    return t


lines = text.split("\n")
meta, i = {}, 0
while i < len(lines) and lines[i].startswith("@@"):
    key, _, val = lines[i][2:].partition(" ")
    meta[key] = val.strip(); i += 1

para(meta["TITLE_EN"], "Heading 1")
para(meta["TITLE_ES"], "Heading 1")
if variant == "named":
    para("Muhammad Albie Fathurrahman", "Autor")
    para("muhamadalbie.2024@student.uny.ac.id", "Email y Universidad")
    para("Universitas Negeri Yogyakarta, Indonesia", "Email y Universidad")
else:
    para("Author", "Autor")
    para("Email", "Email y Universidad")
    para("University", "Email y Universidad")
para("")
para("Date of reception: XXXXX")
para("Date of acceptance: XXXXX")
para("")
labelled("Abstract:", meta["ABSTRACT_EN"])
labelled("Keywords:", meta["KEYWORDS_EN"])
para("")
labelled("Resumen:", meta["ABSTRACT_ES"])
labelled("Palabras clave:", meta["KEYWORDS_ES"])

in_refs = False
buf_table = []
for ln in lines[i:]:
    s = ln.rstrip()
    if s.startswith("|"):
        cells = [c.strip() for c in s.strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            buf_table.append(cells)
        continue
    if buf_table:
        add_table(buf_table); buf_table = []
    if not s:
        continue
    if s.startswith("CAPTION: "):
        para(s[9:], "Tabla título"); continue
    if s.startswith("SOURCE: "):
        para(s[8:], "Tabla fuente"); continue
    if s.startswith("## "):
        para(s[3:], "Heading 2"); in_refs = s[3:].strip() == "References"; continue
    if s.startswith("### "):
        para(s[4:], "Heading 3"); continue
    if s.startswith("- "):
        para(s[2:], "Viñetas/listas"); continue
    para(s, "Referencias" if in_refs else "normal")
if buf_table:
    add_table(buf_table)

# Word comments for the reviewer-notes version
if comments_path:
    notes = json.load(open(comments_path, encoding="utf-8"))
    paras = d.paragraphs
    missing = []
    for n in notes:
        target = next((p for p in paras if n["anchor"] in p.text), None)
        if target is None or not target.runs:
            missing.append(n["anchor"]); continue
        d.add_comment(target.runs, text=n["text"], author=n.get("author", "Reviewer"), initials=n.get("initials", "RV"))
    if missing:
        print("UNANCHORED COMMENTS:", missing)

# Appendix (Indonesian review report and response table)
if appendix_path:
    p = d.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)
    buf_table = []
    for ln in open(appendix_path, encoding="utf-8").read().split("\n"):
        s = ln.rstrip()
        if s.startswith("|"):
            cells = [c.strip().replace("\x00", "|") for c in s.replace("\\|", "\x00").strip("|").split("|")]
            if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                buf_table.append(cells)
            continue
        if buf_table:
            add_table(buf_table); buf_table = []
        if not s or s.startswith("---"):
            continue
        if s.startswith("# "):
            para(s[2:], "Heading 1"); continue
        if s.startswith("## "):
            para(s[3:], "Heading 2"); continue
        if s.startswith("### "):
            para(s[4:], "Heading 3"); continue
        if re.match(r"^\s*[-*] ", s):
            para(re.sub(r"^\s*[-*] ", "", s), "Viñetas/listas"); continue
        if s.startswith(">"):
            para(s.lstrip("> ").strip(), "Citas"); continue
        para(s.replace("`", ""))
    if buf_table:
        add_table(buf_table)

# hyperlinks/URLs plain black: nothing to do (added as plain text)
d.save(out)
print("saved", out)
